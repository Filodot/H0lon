"""Source kind detection: by URL, by file extension and magic bytes, and by a PDF page profile.

PDF heuristics
--------------
Every page (up to `PROFILE_MAX_PAGES`, evenly sampled beyond that) is measured with PyMuPDF:
non-whitespace characters of the text layer, page aspect (width / height, rotation applied)
and the share of the page area covered by raster images. The profile gives the signals that
end up in `quality`: `chars_per_page` (median), `text_layer` (share of pages with at least
`TEXT_PAGE_MIN_CHARS` characters), `aspect` (first page), `image_cover` (median) and the
`producer` / `creator` metadata. The kind is chosen in this order:

1. ``slides`` — the creator or producer names presentation software (`SLIDES_SIGNATURES`:
   Beamer, PowerPoint, Keynote, Impress …). The metadata is the most reliable signal: Beamer
   writes «LaTeX with Beamer class», PowerPoint exports name themselves.
2. ``slides`` — landscape pages (median aspect ≥ `SLIDES_MIN_ASPECT` = 1.3, which covers
   4:3, 16:10 and 16:9) with *some* text (`text_layer` ≥ `SLIDES_MIN_TEXT_LAYER` = 0.2: at
   least every fifth slide has a title or bullets) but little of it (median
   ≤ `SLIDES_MAX_CHARS` = 1200 characters; real slides carry 50–800, a landscape handout or
   article page 2000+).
3. ``pdf-scan`` — median below `SCAN_MAX_CHARS` = 40 characters per page: there is no usable
   text layer (phone photos and scanners give 0; a stamp like «Scanned with CamScanner» or a
   page number stays below 40). Landscape pages without any text (book spreads, whiteboard
   photos) land here too, not in `slides`.
4. ``pdf-text`` — everything else. Mixed documents (some pages scanned) are `pdf-text`: the
   extractor classifies every page anyway and sends scanned ones to the vision agent.

Measured on real course materials (copies, 2026-10-03):

======================================  ======  ==========  =========  ======  =========
File                                    pages   chars/page  text pages aspect  kind
======================================  ======  ==========  =========  ======  =========
Beamer lecture, 16:9 (pdfTeX)           31      218         0.97       1.778   slides
LaTeX lab report, A4 (xdvipdfmx)        31      1582        1.00       0.707   pdf-text
Browser «Save as PDF», A4 (Chromium)    25      696         1.00       0.707   pdf-text
iPhone document scan (Quartz)           9       0           0.00       0.773   pdf-scan
Samsung phone scan                      2       0           0.00       0.708   pdf-scan
======================================  ======  ==========  =========  ======  =========

The thresholds have wide margins on these: the slides are told apart by aspect (and by the
Beamer creator) while carrying 218 characters per page against 696–1582 for A4 documents, and
the scans have no text layer at all. Handwritten notes cannot be told from printed scans without
looking at the pages, so a scanned PDF is `pdf-scan`; use ``--kind handwritten`` for a
manuscript (also for a file that is already added: ingest then changes the kind in place).

Damaged files: MuPDF silently rebuilds a broken cross-reference table (`Document.is_repaired`),
and a truncated download (Telegram, a browser) opens with the right page count but empty
pages at the end. Such a PDF is still accepted, but the profile records `repaired` (and
`truncated` when there is no ``%%EOF`` near the end of the file), `quality.notes` says so
instead of promising that the agent will recognise the empty pages, and ingest prints a
warning.
"""

from __future__ import annotations

import contextlib
import html
import os
import re
import statistics
import sys
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, unquote, urlsplit

from h0lon.sources.models import SourceKind

IS_WINDOWS = sys.platform == "win32"

# ---------------------------------------------------------------- thresholds (see docstring)

TEXT_PAGE_MIN_CHARS = 20  # same rule as the extractor: a page with less text is a scan
SCAN_MAX_CHARS = 40
SLIDES_MIN_ASPECT = 1.3
SLIDES_MIN_TEXT_LAYER = 0.2
SLIDES_MAX_CHARS = 1200
OCR_IMAGE_COVER = 0.9  # text layer on top of full-page images: an OCR'ed scan
PROFILE_MAX_PAGES = 300
SLIDES_SIGNATURES = (
    "beamer",
    "powerpoint",
    "keynote",
    "impress",
    "presentation",
)

# ---------------------------------------------------------------- formats

# File format (what the bytes are) -> extension mapping. The kind follows from the format,
# except for PDF, whose kind comes from the page profile.
EXT_FORMAT: dict[str, str] = {
    ".pdf": "pdf",
    ".pptx": "pptx",
    ".docx": "docx",
    ".md": "md",
    ".markdown": "md",
    ".txt": "md",
    ".tex": "tex",
    ".html": "html",
    ".htm": "html",
    ".xhtml": "html",
    ".jpg": "image",
    ".jpeg": "image",
    ".png": "image",
    ".heic": "image",
    ".mp4": "video",
    ".mkv": "video",
    ".webm": "video",
    ".mov": "video",
    ".mp3": "audio",
    ".m4a": "audio",
    ".wav": "audio",
    ".ogg": "audio",
    ".flac": "audio",
}

FORMAT_KIND: dict[str, SourceKind] = {
    "pptx": "slides",
    "docx": "docx",
    "md": "md",
    "tex": "tex",
    "html": "web",
    "image": "handwritten",
    "video": "video",
    "audio": "audio",
}

# Extension used for the copy when the original has none (or an unknown one).
FORMAT_EXT: dict[str, str] = {
    "pdf": ".pdf",
    "pptx": ".pptx",
    "docx": ".docx",
    "md": ".md",
    "tex": ".tex",
    "html": ".html",
}

# Formats people have that we do not read, with what to do about it.
UNSUPPORTED_HINTS: dict[str, str] = {
    ".doc": "сохраните его в Word как .docx",
    ".rtf": "сохраните его в Word как .docx",
    ".odt": "сохраните его как .docx",
    ".ppt": "сохраните его в PowerPoint как .pptx или PDF",
    ".odp": "сохраните его как .pptx или PDF",
    ".djvu": "сконвертируйте его в PDF",
    ".djv": "сконвертируйте его в PDF",
    ".epub": "сконвертируйте его в PDF",
    ".tif": "сохраните страницы как PNG или JPG либо соберите в PDF",
    ".tiff": "сохраните страницы как PNG или JPG либо соберите в PDF",
    ".webp": "сохраните как PNG или JPG",
    ".bmp": "сохраните как PNG или JPG",
    ".zip": "распакуйте архив и добавьте файлы",
    ".rar": "распакуйте архив и добавьте файлы",
    ".7z": "распакуйте архив и добавьте файлы",
}

# ---------------------------------------------------------------- URLs

URL_RE = re.compile(r"^https?://", re.IGNORECASE)
SCHEME_RE = re.compile(r"^[a-z][a-z0-9+.-]*://", re.IGNORECASE)

# Hosts whose every page is a video (or a player).
VIDEO_HOSTS = ("youtube.com", "youtu.be", "youtube-nocookie.com", "vkvideo.ru", "rutube.ru")
# VK has both videos and everything else; only these paths (or ?z=video…) are videos.
VK_HOSTS = ("vk.com", "vk.ru")
VK_VIDEO_PATH_RE = re.compile(r"^/(video|clip|clips)", re.IGNORECASE)
VK_VIDEO_Z_RE = re.compile(r"^(video|clip)-?\d", re.IGNORECASE)


def is_url(text: str) -> bool:
    return bool(URL_RE.match(text.strip()))


def url_problem(url: str) -> str | None:
    """Why an http(s) link cannot be used (a Russian phrase), or None if it parses."""
    try:
        parts = urlsplit(url.strip())
        host = parts.hostname
        _ = parts.port  # raises ValueError for a non-numeric or out-of-range port
    except ValueError as exc:
        message = str(exc).lower()
        if "port" in message:
            return "неверный номер порта"
        if "ipv6" in message:
            return "неверно записан IPv6-адрес"
        return "ссылка записана с ошибкой"
    if not host:
        return "в ссылке нет адреса сайта"
    if any(ch.isspace() for ch in host):
        return "в адресе сайта есть пробел"
    return None


def _host_is(host: str, domain: str) -> bool:
    return host == domain or host.endswith("." + domain)


def is_video_url(url: str) -> bool:
    parts = urlsplit(url.strip())
    host = (parts.hostname or "").lower().rstrip(".")
    if any(_host_is(host, d) for d in VIDEO_HOSTS):
        return True
    if any(_host_is(host, d) for d in VK_HOSTS):
        if VK_VIDEO_PATH_RE.match(parts.path or ""):
            return True
        z = parse_qs(parts.query).get("z", [])
        return any(VK_VIDEO_Z_RE.match(v) for v in z)
    return False


def classify_url(url: str) -> SourceKind:
    """`video` for video hosts and direct links to media files, `audio` for audio files,
    `web` for everything else (the page is downloaded as HTML)."""
    if is_video_url(url):
        return "video"
    suffix = Path(unquote(urlsplit(url.strip()).path or "")).suffix.lower()
    fmt = EXT_FORMAT.get(suffix)
    if fmt in ("video", "audio"):
        return FORMAT_KIND[fmt]
    return "web"


def url_display(url: str) -> str:
    """Domain + path, readable (percent-escapes decoded): «ru.wikipedia.org/wiki/Случайная…»."""
    parts = urlsplit(url.strip())
    host = (parts.hostname or "").lower()
    if host.startswith("www."):
        host = host[4:]
    path = unquote(parts.path or "").rstrip("/")
    text = host + path
    if parts.query:
        text += "?" + unquote(parts.query)
    return text or url.strip()


# ---------------------------------------------------------------- paths


def os_path(path: Path | str) -> str:
    """Absolute path string that works beyond MAX_PATH on Windows without long path support."""
    s = os.path.abspath(os.fspath(path))
    if IS_WINDOWS and len(s) >= 240 and not s.startswith("\\\\?\\"):
        if s.startswith("\\\\"):
            return "\\\\?\\UNC\\" + s[2:]
        return "\\\\?\\" + s
    return s


# ---------------------------------------------------------------- magic bytes


def _read_head(path: Path, n: int = 8192) -> bytes:
    with open(os_path(path), "rb") as fh:
        return fh.read(n)


def _zip_format(path: Path) -> str | None:
    try:
        with zipfile.ZipFile(os_path(path)) as zf:
            names = set(zf.namelist())
    except (zipfile.BadZipFile, OSError):
        return None
    if "word/document.xml" in names:
        return "docx"
    if "ppt/presentation.xml" in names:
        return "pptx"
    return None


def sniff_format(path: Path) -> str | None:
    """Format from the first bytes of the file; None if it is not recognised."""
    try:
        head = _read_head(path)
    except OSError:
        return None
    if b"%PDF-" in head[:1024]:
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        return _zip_format(path)
    if head.startswith(b"\xff\xd8\xff") or head.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image"
    if len(head) >= 12 and head[4:8] == b"ftyp":
        brand = head[8:12]
        if brand in (b"heic", b"heix", b"mif1", b"msf1", b"heim", b"heis"):
            return "image"
        if brand in (b"M4A ", b"M4B "):
            return "audio"
        return "video"
    if head.startswith(b"\x1a\x45\xdf\xa3"):  # Matroska / WebM
        return "video"
    if (
        head.startswith((b"ID3", b"OggS", b"fLaC"))
        or (head[:2] in (b"\xff\xfb", b"\xff\xf3", b"\xff\xf2"))
        or (head.startswith(b"RIFF") and head[8:12] == b"WAVE")
    ):
        return "audio"
    codec = wide_text_codec(head)
    if codec:  # UTF-16/32 text; ingest stores the copy as UTF-8
        text: str | None = head[: len(head) // 4 * 4].decode(codec, errors="replace")
    else:
        text = _as_text(head)
    if text is None:
        return None
    low = text.lstrip("﻿ \t\r\n").lower()
    if low.startswith(("<!doctype html", "<html", "<?xml")) and "<html" in low:
        return "html"
    if "\\documentclass" in text or "\\begin{document}" in text:
        return "tex"
    return "md"


def _as_text(head: bytes) -> str | None:
    if b"\x00" in head:
        return None
    for encoding in ("utf-8", "cp1251"):
        try:
            return head.decode(encoding)
        except UnicodeDecodeError:
            # A multi-byte character cut at the end of the chunk is not an error.
            try:
                return head[:-3].decode(encoding)
            except UnicodeDecodeError:
                continue
    return None


def wide_text_codec(head: bytes) -> str | None:
    """Codec of a text file stored in UTF-16/UTF-32, from its first bytes.

    Windows tools still write these: PowerShell 5.1 ``>`` and «Unicode» in old Notepad give
    UTF-16 LE with a BOM. Returns "utf-32"/"utf-16" for a BOM, "utf-16-le"/"utf-16-be" when
    NUL bytes sit only at odd/even positions (no BOM), "" when NUL bytes look like binary
    data, and None for an ordinary (UTF-8 or single-byte) text file.
    """
    import codecs

    if head.startswith((codecs.BOM_UTF32_LE, codecs.BOM_UTF32_BE)):
        return "utf-32"
    if head.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return "utf-16"
    if b"\x00" not in head:
        return None
    sample = head[: len(head) // 2 * 2]
    half = len(sample) // 2
    if not half:
        return ""
    even, odd = sample[0::2].count(0), sample[1::2].count(0)
    # ASCII characters (spaces, digits, punctuation) give a NUL in the high byte.
    if even <= half // 100 and odd / half >= 0.02:
        candidate = "utf-16-le"
    elif odd <= half // 100 and even / half >= 0.02:
        candidate = "utf-16-be"
    else:
        return ""
    try:
        text = sample.decode(candidate)
    except UnicodeDecodeError:
        try:  # a surrogate pair cut at the end of the chunk
            text = sample[:-2].decode(candidate)
        except UnicodeDecodeError:
            return ""
    if not text:
        return ""
    printable = sum(1 for ch in text if ch.isprintable() or ch in "\t\r\n")
    return candidate if printable >= 0.95 * len(text) else ""


_SUFFIX_RE = re.compile(r"\.[A-Za-z0-9]{1,8}")


def real_suffix(path: Path | str) -> str:
    """Lower-case extension, or "" when the name has none.

    «Л3. Теория меры» has no extension even though `Path.suffix` returns «. Теория меры».
    """
    suffix = Path(path).suffix
    return suffix.lower() if _SUFFIX_RE.fullmatch(suffix) else ""


def file_format(path: Path) -> str | None:
    """Format by extension, falling back to magic bytes for unknown or missing extensions.

    Plain text (Markdown, LaTeX, HTML without an extension) is recognised only in files
    without any extension: `notes.csv` or `script.py` are not sources.
    """
    suffix = real_suffix(path)
    fmt = EXT_FORMAT.get(suffix)
    if fmt is not None:
        return fmt
    if suffix in UNSUPPORTED_HINTS:
        return None
    sniffed = sniff_format(path)
    if suffix and sniffed in ("md", "tex", "html"):
        return None
    return sniffed


def copy_ext(path: Path, fmt: str) -> str:
    """Extension of the copy in sources/: the original one if it matches the format."""
    suffix = real_suffix(path)
    if EXT_FORMAT.get(suffix) == fmt:
        return suffix
    return FORMAT_EXT.get(fmt, suffix)


# ---------------------------------------------------------------- PDF profile


@dataclass
class PdfProfile:
    pages: int
    sampled: int  # pages actually measured (all, or PROFILE_MAX_PAGES evenly spread)
    chars_per_page: float  # median non-whitespace characters of the text layer
    text_layer: float  # share of measured pages with ≥ TEXT_PAGE_MIN_CHARS characters
    aspect: float  # width / height of the first page (rotation applied)
    aspect_median: float
    image_cover: float  # median share of the page area covered by raster images
    producer: str = ""
    creator: str = ""
    repaired: bool = False  # MuPDF rebuilt a broken file (Document.is_repaired)
    truncated: bool = False  # repaired and no %%EOF near the end: an incomplete download
    broken_pages: int = 0  # measured pages that could not be read at all

    @property
    def damage_note(self) -> str | None:
        """Russian sentence about a damaged file, or None for an intact one."""
        if self.truncated:
            return (
                "PDF обрезан (похоже, недокачан) и был восстановлен при чтении: страницы "
                "в конце могут быть пустыми или неполными — скачайте файл заново."
            )
        if self.repaired or self.broken_pages:
            return (
                "PDF повреждён и был восстановлен при чтении: часть страниц может быть "
                "пустой или неполной — проверьте файл."
            )
        return None

    @property
    def signature(self) -> str:
        return f"{self.creator} {self.producer}".lower()

    @property
    def slides_software(self) -> bool:
        return any(sig in self.signature for sig in SLIDES_SIGNATURES)


class InspectError(Exception):
    """The file cannot be used as a source (corrupt, encrypted, empty); message in Russian."""


def _sample_indices(count: int, limit: int) -> list[int]:
    if count <= limit:
        return list(range(count))
    step = count / limit
    return sorted({int(i * step) for i in range(limit)})


def _has_eof_marker(path: Path) -> bool:
    """True if «%%EOF» is in the last 2 KB (readers look within the last 1 KB; some
    generators append a little junk after it)."""
    try:
        with open(os_path(path), "rb") as fh:
            fh.seek(0, os.SEEK_END)
            size = fh.tell()
            fh.seek(max(0, size - 2048))
            return b"%%EOF" in fh.read()
    except OSError:
        return True


@contextlib.contextmanager
def _quiet_mupdf() -> Iterator[None]:
    """MuPDF prints repair messages («format error: …») to stderr; we report our own."""
    import pymupdf

    try:
        previous = pymupdf.TOOLS.mupdf_display_errors()
        pymupdf.TOOLS.mupdf_display_errors(False)
    except Exception:  # pragma: no cover - older PyMuPDF
        yield
        return
    try:
        yield
    finally:
        with contextlib.suppress(Exception):
            pymupdf.TOOLS.mupdf_display_errors(previous)


def profile_pdf(path: Path) -> PdfProfile:
    with _quiet_mupdf():
        return _profile_pdf(path)


def _profile_pdf(path: Path) -> PdfProfile:
    import pymupdf

    try:
        doc = pymupdf.open(os_path(path), filetype="pdf")
    except Exception as exc:  # MuPDF raises its own types (FileDataError, RuntimeError …)
        raise InspectError(f"не удалось открыть PDF: {exc}") from exc
    with doc:
        if doc.needs_pass:
            raise InspectError("PDF защищён паролем — снимите защиту и добавьте файл снова")
        if doc.page_count == 0:
            if getattr(doc, "is_repaired", False) or not _has_eof_marker(path):
                raise InspectError(
                    "PDF повреждён или обрезан (похоже, недокачан): страниц не найдено — "
                    "скачайте файл заново"
                )
            raise InspectError("в PDF нет страниц")
        meta = doc.metadata or {}
        chars: list[int] = []
        aspects: list[float] = []
        covers: list[float] = []
        first_aspect = 0.0
        broken = 0
        for n, index in enumerate(_sample_indices(doc.page_count, PROFILE_MAX_PAGES)):
            try:
                page = doc[index]
                rect = page.rect
                aspect = rect.width / rect.height if rect.height else 0.0
                text = page.get_text("text")
                cover = 0.0
                area = rect.width * rect.height
                if area > 0:
                    for info in page.get_image_info():
                        box = pymupdf.Rect(info["bbox"]) & rect
                        if not box.is_empty:
                            cover += box.width * box.height
                    cover = min(cover / area, 1.0)
            except Exception:
                # A broken page counts as a page without text.
                text, aspect, cover = "", aspects[-1] if aspects else 0.0, 0.0
                broken += 1
            if n == 0:
                first_aspect = aspect
            chars.append(len("".join(text.split())))
            aspects.append(aspect)
            covers.append(cover)
        # Read after the pages: MuPDF may repair lazily when it meets a broken object.
        repaired = bool(getattr(doc, "is_repaired", False))
        return PdfProfile(
            pages=doc.page_count,
            sampled=len(chars),
            chars_per_page=float(statistics.median(chars)),
            text_layer=sum(c >= TEXT_PAGE_MIN_CHARS for c in chars) / len(chars),
            aspect=first_aspect,
            aspect_median=float(statistics.median(aspects)),
            image_cover=float(statistics.median(covers)),
            producer=(meta.get("producer") or "").strip(),
            creator=(meta.get("creator") or "").strip(),
            repaired=repaired,
            truncated=repaired and not _has_eof_marker(path),
            broken_pages=broken,
        )


def classify_pdf(profile: PdfProfile) -> SourceKind:
    """Kind of a PDF from its profile; the rules and thresholds are in the module docstring."""
    if profile.slides_software:
        return "slides"
    if (
        profile.aspect_median >= SLIDES_MIN_ASPECT
        and profile.text_layer >= SLIDES_MIN_TEXT_LAYER
        and profile.chars_per_page <= SLIDES_MAX_CHARS
    ):
        return "slides"
    if profile.chars_per_page < SCAN_MAX_CHARS:
        return "pdf-scan"
    return "pdf-text"


def pdf_signals(profile: PdfProfile, kind: str) -> tuple[dict[str, int | float], dict[str, Any]]:
    """`units` and ingest-time `quality` for a PDF source."""
    units: dict[str, int | float] = {"pages": profile.pages}
    if kind == "slides":
        units["slides"] = profile.pages
    quality: dict[str, Any] = {
        "text_layer": round(profile.text_layer, 3),
        "chars_per_page": round(profile.chars_per_page),
        "aspect": round(profile.aspect, 3),
        "image_cover": round(profile.image_cover, 3),
    }
    if profile.producer:
        quality["producer"] = profile.producer
    if profile.creator:
        quality["creator"] = profile.creator
    notes: list[str] = []
    damage = profile.damage_note
    if damage:
        if profile.repaired:
            quality["repaired"] = True
        if profile.truncated:
            quality["truncated"] = True
        if profile.broken_pages:
            quality["broken_pages"] = profile.broken_pages
        notes.append(damage)
    if profile.sampled < profile.pages:
        notes.append(f"Профиль PDF построен по {profile.sampled} из {profile.pages} страниц.")
    if damage and kind != "pdf-scan" and profile.text_layer < 0.8:
        # Empty pages of a damaged file are lost content, not scans for the agent.
        with_text = round(profile.text_layer * profile.sampled)
        notes.append(f"Текст есть только на {with_text} из {profile.sampled} страниц.")
    elif kind == "pdf-scan" and profile.text_layer == 0:
        notes.append("Нет текстового слоя: все страницы распознает агент.")
    elif kind == "pdf-scan":
        notes.append(
            "Текстовый слой почти пуст: страницы распознает агент, текст слоя — только подсказка."
        )
    elif profile.text_layer >= 0.5 and profile.image_cover >= OCR_IMAGE_COVER:
        notes.append(
            "Страницы — сканы с распознанным текстовым слоем (OCR): в тексте возможны ошибки."
        )
    elif kind == "pdf-text" and profile.text_layer < 0.8:
        with_text = round(profile.text_layer * profile.sampled)
        notes.append(
            f"Текст есть только на {with_text} из {profile.sampled} страниц: "
            "остальные распознает агент."
        )
    if notes:
        quality["notes"] = notes
    return units, quality


# ---------------------------------------------------------------- PPTX profile


def _shape_texts(shapes: Any) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    out: list[str] = []
    for shape in shapes:
        try:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                out.extend(_shape_texts(shape.shapes))
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                out.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                out.extend(cell.text for row in shape.table.rows for cell in row.cells)
        except Exception:
            continue
    return out


def _notes_text(slide: Any) -> str:
    try:
        if not slide.has_notes_slide:
            return ""
        frame = slide.notes_slide.notes_text_frame
        return frame.text if frame is not None else ""
    except Exception:
        return ""


def pptx_signals(path: Path) -> tuple[dict[str, int | float], dict[str, Any]]:
    """`units` (slides) and `quality` (aspect, chars per slide, text share) of a PPTX file."""
    from pptx import Presentation

    try:
        prs = Presentation(os_path(path))
    except Exception as exc:
        raise InspectError(f"не удалось открыть PPTX: {exc}") from exc
    slides = list(prs.slides)
    if not slides:
        raise InspectError("в презентации нет слайдов")
    chars = [len("".join("".join(_shape_texts(s.shapes)).split())) for s in slides]
    with_notes = sum(1 for s in slides if _notes_text(s).strip())
    width, height = prs.slide_width or 0, prs.slide_height or 0
    quality: dict[str, Any] = {
        "text_layer": round(sum(c >= TEXT_PAGE_MIN_CHARS for c in chars) / len(chars), 3),
        "chars_per_page": round(statistics.median(chars)),
    }
    if width and height:
        quality["aspect"] = round(width / height, 3)
    if with_notes:
        quality["slides_with_notes"] = with_notes
    return {"slides": len(slides)}, quality


# ---------------------------------------------------------------- DOCX / HTML checks


def check_docx(path: Path) -> None:
    if _zip_format(path) != "docx":
        raise InspectError("файл не похож на документ DOCX (повреждён или другой формат)")


_META_CHARSET_RE = re.compile(rb"""<meta[^>]+charset\s*=\s*["']?\s*([A-Za-z0-9._:-]+)""", re.I)


def html_charset(data: bytes) -> str | None:
    """Charset from a BOM or a <meta> declaration in the first 4 KB (validated), else None."""
    import codecs

    if data.startswith(codecs.BOM_UTF8):
        return "utf-8"
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return "utf-16"
    m = _META_CHARSET_RE.search(data[:4096])
    if m:
        name = m.group(1).decode("ascii", "ignore")
        try:
            return codecs.lookup(name).name
        except LookupError:
            return None
    return None


@dataclass
class DecodedHtml:
    text: str
    encoding: str  # codec the text was decoded with
    replaced: int = 0  # undecodable bytes replaced with U+FFFD
    guessed: bool = False  # no usable declaration: the encoding was guessed


# A decoding with replacements is trusted when few of its non-ASCII characters are broken:
# one stray byte in a UTF-8 page gives 1 bad character among thousands, while cp1251 bytes
# read as UTF-8 break almost every Cyrillic letter.
_MAX_BAD_SHARE = 0.05
_FEW_BAD = 3


def _lenient(data: bytes, name: str) -> tuple[str, int, int]:
    """(text, replaced characters, non-ASCII characters) of a decoding with replacement."""
    text = data.decode(name, errors="replace")
    bad = text.count("�")
    if name == "utf-8":
        bad -= data.count("�".encode())  # U+FFFD already in the page
    non_ascii = sum(1 for ch in text if ord(ch) > 127)
    return text, max(bad, 0), non_ascii


def _acceptable(bad: int, non_ascii: int) -> bool:
    share = bad / max(non_ascii, 1)
    return share <= _MAX_BAD_SHARE or (bad <= _FEW_BAD and share < 0.5)


def decode_html_detailed(data: bytes, declared: str | None = None) -> DecodedHtml:
    """Decode HTML bytes. Declared charsets (HTTP header > BOM > <meta>) come first.

    1. The first declared charset that decodes the page strictly wins.
    2. If every declared charset fails, the one with the fewest broken characters is used
       with replacement, provided the damage is small (a stray byte must not turn a UTF-8
       page into cp1251 mojibake). A declaration that breaks most characters is a lie
       (servers do lie) and is ignored.
    3. Without a usable declaration: strict UTF-8, then UTF-8 with replacement if the
       damage is small, then cp1251; the last resort is UTF-8 with replacement.
    """
    import codecs

    declared_names: list[str] = []
    for name in (declared, html_charset(data)):
        if name:
            try:
                norm = codecs.lookup(name).name
            except LookupError:
                continue
            if norm not in declared_names:
                declared_names.append(norm)

    def done(text: str, name: str, bad: int = 0, guessed: bool = False) -> DecodedHtml:
        return DecodedHtml(text.lstrip("﻿"), name, bad, guessed)

    for name in declared_names:
        try:
            return done(data.decode(name), name)
        except (UnicodeDecodeError, LookupError):
            continue
    best: tuple[int, int, str, str] | None = None  # (bad, order, name, text)
    for order, name in enumerate(declared_names):
        try:
            text, bad, non_ascii = _lenient(data, name)
        except LookupError:
            continue
        if _acceptable(bad, non_ascii) and (best is None or (bad, order) < best[:2]):
            best = (bad, order, name, text)
    if best is not None:
        return done(best[3], best[2], best[0])

    guessed = True
    try:
        return done(data.decode("utf-8"), "utf-8", guessed=guessed)
    except UnicodeDecodeError:
        pass
    text, bad, non_ascii = _lenient(data, "utf-8")
    if _acceptable(bad, non_ascii):
        return done(text, "utf-8", bad, guessed)
    try:
        return done(data.decode("cp1251"), "cp1251", guessed=guessed)
    except UnicodeDecodeError:
        pass
    return done(text, "utf-8", bad, guessed)


def decode_html(data: bytes, declared: str | None = None) -> tuple[str, str]:
    """(text, encoding) of HTML bytes; see `decode_html_detailed` for the rules."""
    decoded = decode_html_detailed(data, declared)
    return decoded.text, decoded.encoding


class _TitleParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.title: str | None = None
        self.og_title: str | None = None
        self._in_title = False
        self._parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "title" and self.title is None:
            self._in_title = True
        elif tag == "meta" and self.og_title is None:
            a = {k.lower(): (v or "") for k, v in attrs}
            if a.get("property", "").lower() == "og:title" and a.get("content"):
                self.og_title = a["content"]
        elif tag == "body" and self.title is None and not self._in_title:
            raise _StopParsing

    def handle_endtag(self, tag: str) -> None:
        if tag == "title" and self._in_title:
            self._in_title = False
            self.title = "".join(self._parts)

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self._parts.append(data)


class _StopParsing(Exception):
    pass


def html_title(text: str) -> str | None:
    """Text of the first <title> (or og:title), whitespace collapsed; None if absent."""
    parser = _TitleParser()
    try:
        parser.feed(text[:200_000])
        parser.close()
    except _StopParsing:
        pass
    except Exception:
        pass
    for raw in (parser.title, parser.og_title):
        if raw:
            title = " ".join(html.unescape(raw).split())
            if title:
                return title[:300]
    return None


@dataclass
class Inspection:
    """What a file is: format, kind (detected or given) and ingest-time signals."""

    fmt: str
    kind: SourceKind
    units: dict[str, int | float] = field(default_factory=dict)
    quality: dict[str, Any] = field(default_factory=dict)
    title: str | None = None  # from the content (HTML <title>), if any
    warnings: list[str] = field(default_factory=list)  # accepted, but the user should know


def inspect_file(path: Path, fmt: str, *, kind: str | None = None) -> Inspection:
    """Profile a file of a known format. `kind` overrides the detected kind (signals are
    still measured). Raises InspectError when the file cannot be a source."""
    if fmt == "pdf":
        profile = profile_pdf(path)
        detected = classify_pdf(profile)
        final = kind or detected
        units, quality = pdf_signals(profile, final)
        if kind and kind != detected:
            quality["detected_kind"] = detected
        warnings = [profile.damage_note] if profile.damage_note else []
        return Inspection(fmt, final, units, quality, warnings=warnings)  # type: ignore[arg-type]
    if fmt == "pptx":
        units, quality = pptx_signals(path)
        return Inspection(fmt, kind or "slides", units, quality)  # type: ignore[arg-type]
    if fmt == "docx":
        check_docx(path)
    if fmt == "image":
        return Inspection(fmt, kind or "handwritten", {"pages": 1}, {})  # type: ignore[arg-type]
    title = None
    if fmt == "html":
        try:
            with open(os_path(path), "rb") as fh:
                data = fh.read(2_000_000)
        except OSError as exc:
            raise InspectError(f"не удалось прочитать файл: {exc}") from exc
        title = html_title(decode_html(data)[0])
    return Inspection(fmt, kind or FORMAT_KIND[fmt], title=title)  # type: ignore[arg-type]

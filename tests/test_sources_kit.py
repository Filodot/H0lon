"""Synthetic source files for the ingest tests (no tests here; imported by test_sources_*).

Everything is generated: PDFs with PyMuPDF, PPTX with python-pptx, DOCX as a minimal
OOXML package. User materials never go into the repository.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pymupdf

A4 = (595.0, 842.0)
WIDE = (800.0, 450.0)  # 16:9

LOREM = (
    "Random variables map outcomes to numbers. The distribution function F(t) = P(X <= t) "
    "is non-decreasing and right-continuous. Expectation is the Lebesgue integral of X "
    "with respect to the probability measure. "
)


def _text_of(chars: int) -> str:
    text = ""
    while len("".join(text.split())) < chars:
        text += LOREM
    return text


def make_pdf(
    path: Path,
    *,
    pages: int = 3,
    chars: int = 1500,
    size: tuple[float, float] = A4,
    image: bool = False,
    metadata: dict[str, str] | None = None,
    seed: str = "",
) -> Path:
    """PDF with `pages` pages of about `chars` non-space characters each (0 = no text),
    optionally covered by a raster image (a scan)."""
    doc = pymupdf.open()
    for n in range(pages):
        page = doc.new_page(width=size[0], height=size[1])
        if image:
            pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 80), False)
            pix.clear_with(180 + n % 50)
            page.insert_image(page.rect, pixmap=pix)
        if chars:
            rect = pymupdf.Rect(
                page.rect.x0 + 20, page.rect.y0 + 20, page.rect.x1 - 20, page.rect.y1 - 20
            )
            fontsize = 9 if size[0] > size[1] else 8
            page.insert_textbox(rect, f"{seed} page {n + 1}. " + _text_of(chars), fontsize=fontsize)
    if metadata:
        doc.set_metadata(metadata)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    doc.close()
    return path


def truncate_file(src: Path, dst: Path, fraction: float) -> Path:
    """Copy of `src` cut to `fraction` of its length (an interrupted download)."""
    data = src.read_bytes()
    dst.write_bytes(data[: int(len(data) * fraction)])
    return dst


def break_xref(src: Path, dst: Path) -> Path:
    """Copy of a PDF whose startxref offset is wrong: MuPDF repairs it, %%EOF is intact."""
    data = src.read_bytes()
    m = re.search(rb"startxref\s+(\d+)", data)
    assert m is not None
    dst.write_bytes(data[: m.start(1)] + b"0" * len(m.group(1)) + data[m.end(1) :])
    return dst


def make_pptx(path: Path, *, slides: int = 4, notes: bool = True) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for n in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Slide {n + 1}: distribution functions"
        slide.placeholders[1].text = "Right-continuous\nNon-decreasing\nLimits 0 and 1"
        if notes:
            slide.notes_slide.notes_text_frame.text = "Speaker notes"
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


def make_docx(path: Path, text: str = "Hello") -> Path:
    """Minimal OOXML word document (enough for detection; Pandoc can read it too)."""
    content_types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.'
        'relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.'
        'openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
        'relationships/officeDocument" Target="word/document.xml"/></Relationships>'
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:body></w:document>"
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", rels)
        zf.writestr("word/document.xml", document)
    return path


def make_png(path: Path) -> Path:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 30), False)
    pix.clear_with(200)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(pix.tobytes("png"))
    return path

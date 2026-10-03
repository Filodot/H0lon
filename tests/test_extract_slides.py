"""Slides extractor (h0lon/extract/slides.py): PDF slides and PPTX.

`tests/fixtures/m1/pdf/slides_math.pdf` is compiled by XeLaTeX from the .tex next to it
(454×255 pt, CMSY bullets like Beamer). PPTX decks are built with python-pptx in the tests;
LibreOffice is replaced by a fake converter script. No agents: a fake `run_task`.
"""

from __future__ import annotations

import shutil
import sys
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pymupdf
import pytest

from h0lon.agents import Usage
from h0lon.agents.runner import RunResult
from h0lon.extract import pages as pg
from h0lon.extract import slides as sl
from h0lon.extract import vision
from h0lon.extract.model import ExtractContext
from h0lon.sources.models import SourceRecord

pytestmark = pytest.mark.usefixtures("clean_h0lon_env")

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "m1" / "pdf"
SLIDES_PDF = FIXTURES / "slides_math.pdf"


class FakeAgent:
    def __init__(self, write=None) -> None:
        self.calls: list[dict[str, Any]] = []
        self.write = write or (lambda n: f"Распознано: слайд {n}")

    def __call__(
        self, bundle, *, settings, tier="strong", backend=None, fallback=True, on_event=None
    ):
        names = [f.path for f in bundle.contract.files]
        self.calls.append(
            {
                "files": names,
                "task": bundle.task_path.read_text(encoding="utf-8"),
                "images": [p.name for p in bundle.images],
                "hints": {
                    p.name: p.read_text(encoding="utf-8")
                    for p in bundle.inputs_dir.iterdir()
                    if p.suffix == ".txt"
                },
            }
        )
        for name in names:
            (bundle.out_dir / name).write_text(self.write(int(name[1:5])), encoding="utf-8")
        return RunResult(
            ok=True,
            bundle=bundle,
            backend_used="claude",
            attempts=[],
            usage_total=Usage(),
            final_text="",
            problems=[],
        )


def _ctx(tmp_path: Path, settings: Any, src: Path, **kw: Any) -> ExtractContext:
    topic = tmp_path / "topic"
    (topic / "sources").mkdir(parents=True, exist_ok=True)
    (topic / "runs").mkdir(exist_ok=True)
    target = topic / "sources" / f"S1_deck{src.suffix}"
    shutil.copyfile(src, target)
    rec = SourceRecord(
        id="S1",
        kind="slides",
        title="deck",
        file=f"sources/{target.name}",
        added=datetime.now(UTC).isoformat(),
    )
    out = topic / "extracted" / "S1"
    out.mkdir(parents=True, exist_ok=True)
    return ExtractContext(topic_dir=topic, source=rec, settings=settings, out_dir=out, **kw)


# ---------------------------------------------------------------- PDF slides


def _line(index: int, text: str, bbox: tuple[float, float, float, float], size: float, **kw):
    return pg.LineInfo(
        index=index,
        text=text,
        bbox=bbox,
        size=size,
        font=kw.get("font", "SFSX"),
        bold=False,
        block=kw.get("block", index),
    )


def test_slide_title_largest_top_line_with_continuation() -> None:
    prof = pg.PageProfile(number=1, width=454, height=255)
    prof.lines = [
        _line(0, "Непараметрические", (147, 83, 320, 100), 17.2),
        _line(1, "и метрические методы", (127, 96, 327, 113), 17.2),  # below the top third
        _line(2, "Курс лекций по машинному обучению", (136, 139, 318, 150), 10.9, font="SFSS"),
    ]
    title, lines = sl.slide_title(prof)
    assert title == "Непараметрические и метрические методы"
    assert [ln.index for ln in lines] == [0, 1]


def test_slide_title_same_size_body_is_not_glued() -> None:
    prof = pg.PageProfile(number=4, width=454, height=255)
    prof.lines = [
        _line(0, "Задача классификации (напоминание)", (14, 9, 296, 23), 14.3),
        _line(1, "Дано", (28, 56, 59, 70), 14.3, font="SFSS"),
        _line(2, "4/31", (420, 243, 442, 253), 10.0),
    ]
    assert sl.slide_title(prof)[0] == "Задача классификации (напоминание)"
    empty = pg.PageProfile(number=5, width=454, height=255)
    empty.lines = [_line(0, "5/31", (420, 10, 442, 20), 30.0)]
    assert sl.slide_title(empty) == (None, [])


def test_slide_without_title_keeps_its_list() -> None:
    """A slide of bullets only: no bullet becomes the title, the list stays whole."""
    prof = pg.PageProfile(number=1, width=454, height=255)
    prof.lines = [
        _line(0, "• Первый пункт без заголовка", (30, 30, 300, 42), 12, font="Arial", block=0),
        _line(1, "• Второй пункт про метрики", (30, 60, 300, 72), 12, font="Arial", block=0),
        _line(2, "• Третий пункт про деревья", (30, 90, 300, 102), 12, font="Arial", block=0),
    ]
    assert sl.slide_title(prof) == (None, [])
    md = sl.slide_text_markdown(prof)
    assert md.splitlines() == [
        "- Первый пункт без заголовка",
        "- Второй пункт про метрики",
        "- Третий пункт про деревья",
    ]
    numbered = pg.PageProfile(number=2, width=454, height=255)
    numbered.lines = [
        _line(0, "1. Первый шаг алгоритма", (30, 30, 300, 42), 12, font="Arial", block=0),
        _line(1, "2. Второй шаг алгоритма", (30, 60, 300, 72), 12, font="Arial", block=0),
    ]
    assert sl.slide_title(numbered) == (None, [])
    assert sl.slide_text_markdown(numbered).count("1. ") == 2  # Pandoc renumbers 1., 2.


def test_slide_title_must_stand_out_from_the_body() -> None:
    para = pg.PageProfile(number=1, width=454, height=255)
    para.lines = [
        _line(0, "Обычный абзац текста на слайде", (30, 30, 400, 42), 12, font="Arial"),
        _line(1, "и его продолжение на следующей строке", (30, 44, 400, 56), 12, font="Arial"),
        _line(2, "и ещё одна строка того же абзаца", (30, 58, 400, 70), 12, font="Arial"),
    ]
    assert sl.slide_title(para) == (None, [])
    # Beamer: the frame title SFSX 14.35 over body text SFSS 17.22 (lecture slide 17)
    beamer = pg.PageProfile(number=17, width=454, height=255)
    beamer.lines = [
        _line(0, "Одномерный случай", (14, 9, 140, 23), 14.35, font="SFSX1440"),
        _line(
            1,
            "Если Pr([a, b]) является мерой вероятности, то",
            (28, 51, 430, 68),
            17.22,
            font="SFSS1728",
        ),
        _line(2, "p(x) = lim", (60, 91, 160, 108), 17.22, font="CMMI12"),
        _line(
            3,
            "Эмпирическая оценка вероятности с окном",
            (28, 127, 430, 144),
            17.22,
            font="SFSS1728",
        ),
    ]
    assert sl.slide_title(beamer)[0] == "Одномерный случай"
    section = pg.PageProfile(number=3, width=454, height=255)
    section.lines = [
        _line(0, "2. Метрические методы", (100, 60, 350, 80), 20.0, font="Arial-Bold"),
        _line(1, "1. kNN", (40, 120, 200, 132), 12, font="Arial"),
        _line(2, "2. Парзеновское окно", (40, 140, 200, 152), 12, font="Arial"),
    ]
    assert sl.slide_title(section)[0] == "2. Метрические методы"


def test_slide_titles_in_the_top_band_are_kept(tmp_path: Path, settings: Any) -> None:
    """«Пример 1»…«Пример 6» sit in the top 8 % but are titles, not a running header."""
    doc = pymupdf.open()
    for i in range(1, 7):
        page = doc.new_page(width=454, height=255)
        # within the top band (8 % of 255 pt = 20.4 pt); Cyrillic needs the HTML fonts
        page.insert_htmlbox(
            pymupdf.Rect(20, 0, 440, 20), f'<p style="font-size:14pt">Пример {i}</p>'
        )
        page.insert_htmlbox(pymupdf.Rect(30, 90, 440, 120), f"• содержимое слайда {i}")
    pdf = tmp_path / "band.pdf"
    doc.save(str(pdf))
    out = sl.SlidesExtractor().extract(_ctx(tmp_path, settings, pdf, use_vision=False))
    body = out.body_md.read_text(encoding="utf-8")
    assert all(f"## [[S1:s{i}]] Слайд {i}. Пример {i}" in body for i in range(1, 7))
    assert out.quality["slides_titled"] == 6


def test_slide_with_a_small_picture_goes_to_the_agent(tmp_path: Path) -> None:
    """Text slides keep no pictures: a unique picture of ~5 % is content for the agent."""
    import os

    doc = pymupdf.open()
    for i in range(2):
        page = doc.new_page(width=454, height=255)
        page.insert_htmlbox(pymupdf.Rect(14, 9, 400, 30), f"<b>Слайд {i + 1}</b>")
        page.insert_htmlbox(pymupdf.Rect(20, 60, 300, 200), "Обычный текст слайда без формул.")
        noise = pymupdf.Pixmap(pymupdf.csRGB, 40, 40, os.urandom(40 * 40 * 3), False)
        if i == 0:
            page.insert_image(pymupdf.Rect(320, 80, 395, 155), pixmap=noise)
    pdf = tmp_path / "pic.pdf"
    doc.save(str(pdf))
    with pg.opened(pdf) as d:
        first, second = pg.profile_document(d, slides=True)
    assert first.pictures == 1 and (first.kind, first.reason) == ("graphic", "graphic")
    assert second.pictures == 0 and second.kind == "text"


def test_slide_text_lists_and_levels() -> None:
    prof = pg.PageProfile(number=1, width=454, height=255)
    prof.lines = [
        _line(0, "Дано", (28, 40, 60, 52), 12, block=0),
        _line(1, "∙Первый пункт, который", (38, 60, 300, 72), 12, block=1),
        _line(2, "продолжается на второй строке", (50, 74, 300, 86), 12, block=1),
        _line(3, "– вложенный пункт", (55, 90, 300, 102), 12, block=1),
        _line(4, "∙Второй пункт", (38, 106, 300, 118), 12, block=1),
        _line(5, "1. нумерованный", (38, 130, 300, 142), 12, block=2),
        _line(6, "Итог: всё*важно", (28, 160, 300, 172), 12, block=3),
    ]
    md = sl.slide_text_markdown(prof)
    assert md == (
        "Дано\n\n"
        "- Первый пункт, который продолжается на второй строке\n"
        "    - вложенный пункт\n"
        "- Второй пункт\n"
        "- нумерованный".replace("- нумерованный", "1. нумерованный")
        + "\n\nИтог: всё\\*важно"
    )


def test_strip_title_and_heading() -> None:
    assert (
        sl.strip_title("## Расстояние Минковского\n\n$$x$$", "Расстояние  минковского") == "$$x$$"
    )
    assert sl.strip_title("**Итоги**\n\nТекст", "Итоги") == "Текст"
    assert sl.strip_title("Другое начало", "Итоги") == "Другое начало"
    assert sl.slide_heading("S1", 3, "Метод *k* [NN] #1") == (
        "## [[S1:s3]] Слайд 3. Метод \\*k\\* \\[NN\\] \\#1"
    )
    assert sl.slide_heading("S1", 4, None) == "## [[S1:s4]] Слайд 4"


def test_pdf_slides_plan_and_extract_without_vision(tmp_path: Path, settings: Any) -> None:
    ctx = _ctx(tmp_path, settings, SLIDES_PDF)
    plan = sl.SlidesExtractor().plan(ctx)
    assert (plan.pages_total, plan.pages_vision, plan.agent_runs) == (3, 1, 1)
    ctx.use_vision = False
    out = sl.SlidesExtractor().extract(ctx)
    body = out.body_md.read_text(encoding="utf-8")
    assert body.startswith("## [[S1:s1]] Слайд 1. Метрические методы\n\n- Объекты одного класса")
    assert "- Классификация по ближайшим соседям." in body
    assert "## [[S1:s2]] Слайд 2. Расстояние Минковского" in body
    assert "## [[S1:s3]] Слайд 3. Ядерное сглаживание и непараметрическая регрессия" in body
    assert "Оценка плотности строится по окну вокруг точки." in body
    assert "/3" not in body  # slide numbers are dropped
    slide2 = body.split("[[S1:s2]]")[1].split("[[S1:s3]]")[0]
    assert '::: {.uncertain title="не распознано: агент отключён"}' in slide2
    assert out.quality["slides_titled"] == 3 and out.quality["pages_math"] == 1
    assert not list(ctx.out_dir.glob("pages/*.png"))  # renders are only for the agent


def test_pdf_slides_with_vision(tmp_path: Path, settings: Any, monkeypatch) -> None:
    agent = FakeAgent(write=lambda n: "### Расстояние Минковского\n\n$$\\rho(a,b)$$")
    monkeypatch.setattr(vision, "run_task", agent)
    ctx = _ctx(tmp_path, settings, SLIDES_PDF)
    out = sl.SlidesExtractor().extract(ctx)
    body = out.body_md.read_text(encoding="utf-8")
    (call,) = agent.calls
    assert call["files"] == ["p0002.md"] and "Страницы — слайды презентации" in call["task"]
    assert "Расстояние Минковского" not in call["hints"]["p0002.txt"]  # title is in the heading
    assert "## [[S1:s2]] Слайд 2. Расстояние Минковского\n\n$$\\rho(a,b)$$" in body
    assert out.agent_runs == 1 and out.pages_vision == 1
    png = pymupdf.Pixmap(str(ctx.out_dir / "pages" / "p0002.png"))
    assert max(png.width, png.height) == pytest.approx(1400, abs=2)


def test_slides_batches_of_ten(tmp_path: Path, settings: Any, monkeypatch) -> None:
    agent = FakeAgent()
    monkeypatch.setattr(vision, "run_task", agent)
    doc = pymupdf.open()
    for i in range(12):
        page = doc.new_page(width=454, height=255)
        page.insert_htmlbox(pymupdf.Rect(20, 10, 440, 40), f"<b>Слайд номер {i + 1}</b>")
        page.insert_htmlbox(
            pymupdf.Rect(20, 60, 440, 200), "Пусть μ ∈ ℝ, ∀ε > 0 ∃δ: |x − y| < δ ⟹ ∑ ≤ ε"
        )
    pdf = tmp_path / "many.pdf"
    doc.save(str(pdf))
    ctx = _ctx(tmp_path, settings, pdf)
    out = sl.SlidesExtractor().extract(ctx)
    # batches run in parallel (agents.parallel_runs = 2): order of calls is not fixed
    assert sorted(len(c["files"]) for c in agent.calls) == [2, 10]
    assert out.agent_runs == 2 and out.pages_vision == 12


def test_pdf_slide_with_table(tmp_path: Path, settings: Any) -> None:
    doc = pymupdf.open()
    page = doc.new_page(width=454, height=255)
    page.insert_htmlbox(pymupdf.Rect(14, 9, 400, 30), "<b>Сравнение методов</b>")
    xs, ys = [40, 160, 280, 400], [60, 90, 120, 150]
    shape = page.new_shape()
    for x in xs:
        shape.draw_line((x, ys[0]), (x, ys[-1]))
    for y in ys:
        shape.draw_line((xs[0], y), (xs[-1], y))
    shape.finish(color=(0, 0, 0), width=0.6)
    shape.commit()
    cells = [["Метод", "Точность", "Время"], ["kNN", "0,93", "1 с"], ["SVM", "0,95", "10 с"]]
    for r, row in enumerate(cells):
        for c, text in enumerate(row):
            page.insert_htmlbox(
                pymupdf.Rect(xs[c] + 4, ys[r] + 6, xs[c + 1] - 4, ys[r + 1] - 2), text
            )
    pdf = tmp_path / "table.pdf"
    doc.save(str(pdf))
    ctx = _ctx(tmp_path, settings, pdf, use_vision=False)
    body = sl.SlidesExtractor().extract(ctx).body_md.read_text(encoding="utf-8")
    assert body.startswith("## [[S1:s1]] Слайд 1. Сравнение методов\n\n|Метод|Точность|Время|")
    assert "|SVM|0,95|10 с|" in body
    assert "### Сравнение методов" not in body  # the title is not repeated in the body


# ---------------------------------------------------------------- PPTX


def _png(path: Path) -> Path:
    import os

    pix = pymupdf.Pixmap(pymupdf.csRGB, 160, 120, os.urandom(160 * 120 * 3), False)
    pix.save(str(path))
    return path


def build_pptx(path: Path, image: Path) -> Path:
    from lxml import etree
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])  # title and content
    s1.shapes.title.text = "Введение"
    tf = s1.placeholders[1].text_frame
    tf.text = "Пункт A"
    p = tf.add_paragraph()
    p.text = "Подпункт"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Пункт B"
    s1.notes_slide.notes_text_frame.text = "Сказать про X.\nВторая строка"

    s2 = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    s2.shapes.title.text = "Таблица"
    table = s2.shapes.add_table(2, 3, Inches(1), Inches(2), Inches(6), Inches(1)).table
    for c, text in enumerate(["Метод", "Точность", "a|b"]):
        table.cell(0, c).text = text
    for c, text in enumerate(["kNN", "0,93", "—"]):
        table.cell(1, c).text = text

    s3 = prs.slides.add_slide(prs.slide_layouts[6])  # blank: a picture only
    s3.shapes.add_picture(str(image), Inches(1), Inches(1), Inches(6))

    s4 = prs.slides.add_slide(prs.slide_layouts[6])
    box = s4.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2)).text_frame
    run = box.paragraphs[0].add_run()
    run.text = "Жирный"
    run.font.bold = True
    run = box.paragraphs[0].add_run()
    run.text = " и обычный текст"
    step = box.add_paragraph()
    step.text = "первый шаг"
    num = etree.SubElement(step._p.get_or_add_pPr(), qn("a:buAutoNum"))
    num.set("type", "arabicPeriod")
    data = CategoryChartData()
    data.categories = ["Янв", "Фев"]
    data.add_series("Продажи", (1.5, 2.0))
    s4.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3.5), Inches(5), Inches(3), data
    )
    prs.save(str(path))
    return path


@pytest.fixture
def deck(tmp_path: Path) -> Path:
    return build_pptx(tmp_path / "deck.pptx", _png(tmp_path / "pic.png"))


def test_read_pptx(tmp_path: Path, deck: Path) -> None:
    slides = sl.read_pptx(deck, figures_dir=tmp_path / "figures", prefix="S1_")
    assert [s.title for s in slides] == ["Введение", "Таблица", None, None]
    assert slides[0].blocks == ["- Пункт A\n    - Подпункт\n- Пункт B"]
    assert slides[0].notes == "Сказать про X.\nВторая строка"
    assert slides[1].blocks == ["| Метод | Точность | a\\|b |\n|---|---|---|\n| kNN | 0,93 | — |"]
    assert not slides[2].has_text and slides[2].pictures == [
        tmp_path / "figures" / "S1_s0003_1.png"
    ]
    assert slides[3].blocks[0] == "**Жирный** и обычный текст"
    assert slides[3].blocks[1] == "1. первый шаг"
    assert slides[3].blocks[2].startswith("::: {.figure-description}\nДиаграмма")
    assert "| Янв | 1.5 |" in slides[3].blocks[2]


def test_pptx_extract_without_libreoffice(
    tmp_path: Path, settings: Any, deck: Path, monkeypatch
) -> None:
    agent = FakeAgent(write=lambda n: "::: {.figure-description}\nШум из цветных точек.\n:::")
    monkeypatch.setattr(vision, "run_task", agent)
    monkeypatch.setattr(sl, "find_soffice", lambda: None)
    ctx = _ctx(tmp_path, settings, deck)
    plan = sl.SlidesExtractor().plan(ctx)
    assert (plan.pages_total, plan.pages_vision, plan.agent_runs) == (4, 1, 1)
    assert sl.NO_SOFFICE_NOTE in plan.notes
    out = sl.SlidesExtractor().extract(ctx)
    body = out.body_md.read_text(encoding="utf-8")
    assert "## [[S1:s1]] Слайд 1. Введение\n\n- Пункт A\n    - Подпункт\n- Пункт B" in body
    assert "Заметки докладчика: Сказать про X.\\\nВторая строка" in body
    assert (
        "## [[S1:s3]] Слайд 3\n\n![](figures/S1_s0003_1.png)\n\n::: {.figure-description}" in body
    )
    (call,) = agent.calls
    assert call["files"] == ["p0003.md"] and call["images"] == ["p0003.png"]
    q = out.quality
    assert sl.NO_SOFFICE_NOTE in q["notes"]
    assert q["slides_with_notes"] == 1 and q["pictures"] == 1 and q["pages_vision"] == 1
    assert q["renderer"] == "python-pptx"


def test_pptx_picture_slide_without_agent(
    tmp_path: Path, settings: Any, deck: Path, monkeypatch
) -> None:
    monkeypatch.setattr(sl, "find_soffice", lambda: None)
    ctx = _ctx(tmp_path, settings, deck, use_vision=False)
    body = sl.SlidesExtractor().extract(ctx).body_md.read_text(encoding="utf-8")
    slide3 = body.split("[[S1:s3]]")[1].split("[[S1:s4]]")[0]
    assert "![](figures/S1_s0003_1.png)" in slide3
    assert '::: {.uncertain title="не распознано: агент отключён"}' in slide3
    assert not list(ctx.out_dir.glob("pages/*.png"))  # nothing prepared for an absent agent


FAKE_SOFFICE = """
import sys
from pathlib import Path
import pymupdf
args = sys.argv[1:]
outdir = Path(args[args.index("--outdir") + 1])
src = Path(args[-1])
picture = str(Path(sys.argv[0]).with_name("pic.png"))
doc = pymupdf.open()
for i in range(4):
    page = doc.new_page(width=720, height=540)
    if i == 2:
        page.insert_image(pymupdf.Rect(60, 60, 660, 480), filename=picture)
    else:
        text = f"<p>Текст слайда {i + 1} без формул и картинок, обычные слова.</p>"
        page.insert_htmlbox(pymupdf.Rect(40, 40, 680, 500), text)
doc.save(str(outdir / (src.stem + ".pdf")))
"""


def test_pptx_with_fake_libreoffice(tmp_path: Path, settings: Any, deck: Path, monkeypatch) -> None:
    script = tmp_path / "fake_soffice.py"
    script.write_text(FAKE_SOFFICE, encoding="utf-8")
    monkeypatch.setattr(sl, "find_soffice", lambda: [sys.executable, str(script)])
    agent = FakeAgent()
    monkeypatch.setattr(vision, "run_task", agent)
    ctx = _ctx(tmp_path, settings, deck)
    out = sl.SlidesExtractor().extract(ctx)
    (call,) = agent.calls
    assert call["files"] == ["p0003.md"]  # the picture slide, rendered from the PDF
    assert (ctx.out_dir / "slides.pdf").is_file()
    assert not (ctx.out_dir / ".soffice").exists()
    body = out.body_md.read_text(encoding="utf-8")
    # the rendered slide is transcribed as a whole, its picture stays linked after it
    assert "## [[S1:s3]] Слайд 3\n\nРаспознано: слайд 3\n\n![](figures/S1_s0003_1.png)" in body
    assert out.quality["renderer"] == "libreoffice"
    assert sl.NO_SOFFICE_NOTE not in out.quality["notes"]


def test_convert_to_pdf_reports_failure(tmp_path: Path) -> None:
    script = tmp_path / "broken.py"
    script.write_text("import sys; sys.stderr.write('boom'); sys.exit(3)", encoding="utf-8")
    pdf, error = sl.convert_to_pdf(
        [sys.executable, str(script)], tmp_path / "x.pptx", tmp_path / "w"
    )
    assert pdf is None and "boom" in error


def test_slide_heading_plain_math_letters() -> None:
    from h0lon.extract.slides import plain_math_letters, slide_heading

    title = "Метод 𝑘-ближайших соседей (𝑘NN)"
    assert plain_math_letters(title) == "Метод k-ближайших соседей (kNN)"
    assert plain_math_letters("x² и 𝐱") == "x² и x"  # superscripts are kept
    assert slide_heading("S1", 13, "Метод 𝑘NN") == "## [[S1:s13]] Слайд 13. Метод kNN"
